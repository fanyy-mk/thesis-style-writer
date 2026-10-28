#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把论文 Markdown 转成基础排版的 Word 文档（宋体/黑体、小四、1.5 倍行距、首行缩进）。

用法：
  python md2docx.py final.md final.docx --title "论文题目"
依赖：python-docx（pip install python-docx）
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError:  # pragma: no cover
    print("缺少依赖 python-docx。请先执行：pip install python-docx")
    raise SystemExit(2)

INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
TABLE_SEP_RE = re.compile(r"^\s*\|?[\s:|-]+\|?\s*$")


def set_run_font(run, cn_font: str = "宋体", size: float = 12, bold: bool = False,
                 italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), cn_font)


def add_inline(paragraph, text: str, cn_font: str = "宋体", size: float = 12,
               bold: bool = False) -> None:
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    parts = INLINE_RE.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, cn_font, size, bold=True)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, "宋体", size - 1, bold=bold)
            run.font.name = "Consolas"
        else:
            run = paragraph.add_run(part)
            set_run_font(run, cn_font, size, bold=bold)


def configure_document(doc) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    try:
        normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    except Exception:  # noqa: BLE001
        pass
    fmt = normal.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)


def add_body(doc, text: str, first_indent: bool = True, size: float = 12) -> None:
    paragraph = doc.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    if first_indent:
        fmt.first_line_indent = Pt(size * 2)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(paragraph, text, size=size)


def add_heading(doc, text: str, level: int, page_break: bool = False) -> None:
    paragraph = doc.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.5
    fmt.first_line_indent = Pt(0)
    fmt.space_before = Pt(18 if level == 1 else 12)
    fmt.space_after = Pt(10 if level == 1 else 6)
    if page_break:
        fmt.page_break_before = True
    if level == 1:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(paragraph, text, cn_font="黑体", size=16, bold=True)
    elif level == 2:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_inline(paragraph, text, cn_font="黑体", size=14, bold=True)
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_inline(paragraph, text, cn_font="黑体", size=12, bold=True)


def add_list_item(doc, text: str, numbered: bool) -> None:
    style = "List Number" if numbered else "List Bullet"
    try:
        paragraph = doc.add_paragraph(style=style)
    except KeyError:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.74)
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = Pt(0)
    add_inline(paragraph, text, size=12)


def add_table(doc, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_idx, row in enumerate(rows):
        for col_idx in range(cols):
            cell = table.cell(row_idx, col_idx)
            text = row[col_idx] if col_idx < len(row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            add_inline(paragraph, text, size=10.5, bold=(row_idx == 0))
    doc.add_paragraph()


def convert(md_text: str, title: str | None, page_break: bool) -> tuple[object, dict]:
    doc = Document()
    configure_document(doc)

    if title:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt = paragraph.paragraph_format
        fmt.space_before = Pt(12)
        fmt.space_after = Pt(18)
        add_inline(paragraph, title, cn_font="黑体", size=22, bold=True)

    lines = md_text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    stats = {"headings": 0, "paragraphs": 0, "tables": 0, "lists": 0, "images_skipped": 0}
    buffer: list[str] = []
    heading_seen = False

    def flush() -> None:
        nonlocal buffer
        text = "".join(buffer).strip()
        buffer = []
        if text:
            add_body(doc, text)
            stats["paragraphs"] += 1

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            flush()
            i += 1
            continue

        if stripped.startswith("```"):
            flush()
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                i += 1
            i += 1
            continue

        if stripped.startswith("!["):
            flush()
            stats["images_skipped"] += 1
            i += 1
            continue

        if re.match(r"^\s*(?:\*\s*){3,}$|^\s*(?:-\s*){3,}$|^\s*(?:_\s*){3,}$", stripped):
            flush()
            i += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            flush()
            level = min(len(heading.group(1)), 3)
            add_heading(doc, heading.group(2).strip(), level,
                        page_break=page_break and heading_seen and level == 1)
            heading_seen = True
            stats["headings"] += 1
            i += 1
            continue

        if stripped.startswith("|"):
            flush()
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = []
            for idx, table_line in enumerate(table_lines):
                if idx == 1 and TABLE_SEP_RE.match(table_line):
                    continue
                cells = [cell.strip() for cell in table_line.strip("|").split("|")]
                rows.append(cells)
            add_table(doc, rows)
            stats["tables"] += 1
            continue

        bullet = re.match(r"^\s*[-*+]\s+(.*)$", line)
        if bullet:
            flush()
            add_list_item(doc, bullet.group(1).strip(), numbered=False)
            stats["lists"] += 1
            i += 1
            continue

        numbered = re.match(r"^\s*\d+[.)]\s+(.*)$", line)
        if numbered:
            flush()
            add_list_item(doc, numbered.group(1).strip(), numbered=True)
            stats["lists"] += 1
            i += 1
            continue

        if stripped.startswith(">"):
            stripped = stripped.lstrip(">").strip()
            if not stripped:
                flush()
                i += 1
                continue
            buffer.append(stripped)
            i += 1
            continue

        buffer.append(stripped)
        i += 1

    flush()
    return doc, stats


def main() -> int:
    parser = argparse.ArgumentParser(description="论文 Markdown → Word")
    parser.add_argument("source", help="输入 Markdown 文件")
    parser.add_argument("target", help="输出 .docx 文件")
    parser.add_argument("--title", help="在文首插入居中标题")
    parser.add_argument("--no-page-break", action="store_true", help="一级标题前不分页")
    args = parser.parse_args()

    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    source = Path(args.source)
    if not source.exists():
        print(f"源文件不存在：{source}")
        return 2

    md_text = source.read_text(encoding="utf-8-sig", errors="ignore")
    doc, stats = convert(md_text, args.title, page_break=not args.no_page_break)

    target = Path(args.target)
    target.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(target)
    except Exception as exc:  # noqa: BLE001
        print(f"保存 Word 失败：{exc}")
        print("请关闭正在占用该文件的 Word 进程后重试。")
        return 2

    print(f"已生成：{target.resolve()}")
    print(f"标题 {stats['headings']} / 正文段 {stats['paragraphs']} / "
          f"表格 {stats['tables']} / 列表 {stats['lists']}")
    if stats["images_skipped"]:
        print(f"[提示] 跳过 {stats['images_skipped']} 张图片；图片与题注请按学校模板手工插入。")
    print("提示：封面、目录、页眉页脚、页码、三线表样式请套用学校模板。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
